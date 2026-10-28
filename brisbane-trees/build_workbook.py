"""Build the Brisbane tree maintenance workbook from Brisbane City Council open data.

Run:  python3 build_workbook.py [output.xlsx]

Re-running refreshes the Council data (Tree_Register reference columns and
Protected_Areas) but starts the Health_Assessments and Maintenance_Log sheets
from scratch, so copy your own log rows across before overwriting a file you
have been filling in.
"""

import datetime
import json
import sys
import urllib.parse
import urllib.request

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ARCGIS = "https://services2.arcgis.com/dEKgZETqwmDAh1rP/arcgis/rest/services"
SHERWOOD = "botanic_collection_sherwood_arboretum"
NALL_LAYERS = [
    ("NALL-CV", "Council Vegetation", "Protected_Vegetation_Natural_Assets_Local_Law_2003_Council_Vegetation"),
    ("NALL-SUV", "Significant Urban Vegetation", "Protected_Vegetation_Natural_Assets_Local_Law_2003_Significant_Urban_Vegetation"),
    ("NALL-SNV", "Significant Native Vegetation", "Protected_Vegetation_Natural_Assets_Local_Law_2003_Significant_Native_Vegetation"),
    ("NALL-WWV", "Waterway and Wetland Vegetation", "Protected_Vegetation_Natural_Assets_Local_Law_2003_Waterway_and_Wetland_Vegetation"),
]
SLT = ("CP-SLT", "City Plan 2014 Significant landscape tree overlay", "Significant_landscape_tree_overlay")

LOG_ROWS = 5000  # rows covered by validations and register formulas
LAST = LOG_ROWS + 1

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
FORMULA_FILL = PatternFill("solid", fgColor="E2EFDA")
EXAMPLE_FONT = Font(name=FONT, italic=True, color="7F7F7F")

# Council's published request reasons and assessment outcomes:
# https://www.brisbane.qld.gov.au/environment-and-water/plants--trees-and-gardening/brisbane-trees
# https://www.brisbane.qld.gov.au/online-services/report-an-issue/report-it-tree-removal
ISSUE_TYPES = [
    "None observed", "Damaged or split", "Leaning", "Dead or dying", "Diseased or unhealthy",
    "Pests", "Fallen or broken branches", "Obstructing footpath, road or sign",
    "Imminent risk", "Storm damage", "Other",
]
OUTCOMES = ["No action", "Pruning", "Health management (tree rehabilitation)", "Removal and replacement planting"]
# Condition and risk scales are this workbook's own; Council publishes neither.
CONDITIONS = ["Good", "Fair", "Poor", "Dead", "Unknown"]
RISKS = ["Low", "Medium", "High", "Extreme"]
WORK_TYPES = [
    "Structural pruning (young tree)", "Clearance pruning", "Dead or diseased wood removal",
    "Pest and disease management", "Storm damage clean-up", "Health management (rehabilitation)",
    "Removal", "Stump grinding", "Replacement planting", "Inspection only",
]
STATUSES = ["Requested", "Scheduled", "In progress", "Completed", "Cancelled"]
CLEARANCES = [
    ("Verge or footpath", 2.4),
    ("Residential or industrial road", 4.5),
    ("Bikeway", 4.5),
    ("Arterial road", 5.0),
    ("Not applicable", None),
]


def get_json(url, params):
    with urllib.request.urlopen(url + "?" + urllib.parse.urlencode(params), timeout=120) as r:
        data = json.load(r)
    if "error" in data:
        raise RuntimeError(f"{url}: {data['error']}")
    return data


def query_all(service, **params):
    """Page through every feature of layer 0 of a hosted Feature Service."""
    url = f"{ARCGIS}/{service}/FeatureServer/0/query"
    base = {"where": "1=1", "outFields": "*", "f": "json", "resultRecordCount": 2000}
    base.update(params)
    feats, offset = [], 0
    while True:
        data = get_json(url, {**base, "resultOffset": offset})
        feats += data["features"]
        if not data.get("exceededTransferLimit") and len(data["features"]) < base["resultRecordCount"]:
            return feats
        offset += len(data["features"])


def point_in_rings(x, y, rings):
    """Even-odd rule over all rings, which handles Esri holes."""
    inside = False
    for ring in rings:
        for (x1, y1), (x2, y2) in zip(ring, ring[1:]):
            if (y1 > y) != (y2 > y) and x < (x2 - x1) * (y - y1) / (y2 - y1) + x1:
                inside = not inside
    return inside


def overlay_hits(trees, service):
    """Return {tree_id: [OBJECTID, ...]} for polygons of `service` containing each tree."""
    xs = [t["longitude"] for t in trees]
    ys = [t["latitude"] for t in trees]
    env = {"xmin": min(xs), "ymin": min(ys), "xmax": max(xs), "ymax": max(ys), "spatialReference": {"wkid": 4326}}
    polys = query_all(service, outFields="OBJECTID", geometry=json.dumps(env), geometryType="esriGeometryEnvelope",
                      inSR=4326, spatialRel="esriSpatialRelIntersects", outSR=4326, returnGeometry="true")
    hits = {}
    for t in trees:
        for p in polys:
            if point_in_rings(t["longitude"], t["latitude"], p["geometry"]["rings"]):
                hits.setdefault(t["Tree_ID"], []).append(p["attributes"]["OBJECTID"])
    return hits


def style_header(ws, headers, widths, fills=None):
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = (fills or {}).get(col, HEADER_FILL)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"


def add_list_validation(ws, col, source, prompt):
    dv = DataValidation(type="list", formula1=source, allow_blank=True, showErrorMessage=True,
                        errorTitle="Not in list", error="Pick a value from the list (see the Lookups sheet).",
                        promptTitle="Pick from list", prompt=prompt, showInputMessage=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{LAST}")


def build(out_path):
    today = datetime.date.today()

    trees = [f["attributes"] for f in query_all(SHERWOOD, returnGeometry="false")]
    trees.sort(key=lambda t: t["Tree_ID"])
    hits = {code: overlay_hits(trees, svc) for code, _, svc in NALL_LAYERS + [SLT]}

    areas = []
    for code, name, svc in NALL_LAYERS:
        for f in query_all(svc, returnGeometry="false", returnCentroid="true", outSR=4326):
            a = f["attributes"]
            areas.append((f"{code}-{a['OBJECTID']}", name, a.get("CATEGORY") or name, None,
                          f["centroid"]["y"], f["centroid"]["x"], a["Shape__Area"]))
    for f in query_all(SLT[2], returnGeometry="false", returnCentroid="true", outSR=4326):
        a = f["attributes"]
        areas.append((f"CP-SLT-{a['OBJECTID']}", SLT[1], a["OVL2_DESC"], a.get("DESCRIPTION"),
                       f["centroid"]["y"], f["centroid"]["x"], a["Shape__Area"]))

    wb = Workbook()
    readme = wb.active
    readme.title = "README"

    # ---- Tree_Register ---------------------------------------------------
    ws = wb.create_sheet("Tree_Register")
    headers = ["Tree_ID", "Source", "Source_Tree_ID", "Latitude", "Longitude", "Scientific_Name", "Common_Name",
               "Family", "Height_m", "Crown_Width_m", "DBH_mm", "Year_Established", "Protected_Vegetation_Areas",
               "In_Significant_Landscape_Tree_Overlay",
               "Assessments", "Last_Assessed", "Latest_Condition", "Latest_Risk", "Latest_Recommended_Outcome",
               "Maintenance_Jobs", "Last_Maintenance", "Open_Jobs"]
    widths = [12, 22, 10, 11, 11, 30, 26, 18, 9, 10, 9, 10, 26, 14, 11, 12, 13, 11, 30, 12, 13, 9]
    style_header(ws, headers, widths, {c: PatternFill("solid", fgColor="548235") for c in range(15, 23)})
    ha, ml = "Health_Assessments", "Maintenance_Log"
    ha_id, ha_date = f"{ha}!$B$2:$B${LAST}", f"{ha}!$C$2:$C${LAST}"
    ml_id, ml_date, ml_status = f"{ml}!$B$2:$B${LAST}", f"{ml}!$D$2:$D${LAST}", f"{ml}!$H$2:$H${LAST}"
    for r, t in enumerate(trees, 2):
        tid = t["Tree_ID"]
        protected = [f"{code}-{oid}" for code, _, _ in NALL_LAYERS for oid in hits[code].get(tid, [])]
        height = t.get("Height")
        try:
            height = float(height) if height not in (None, "") else None
        except ValueError:
            pass  # keep Council's text, e.g. a range
        row = [f"SHW-{tid:04d}", "Sherwood Arboretum", tid, t["latitude"], t["longitude"], t["Scientific_Name"],
               t["Common_name"], t["Family"], height, t.get("Crown_Width"), t.get("DBH"), t.get("Year_Established"),
               ", ".join(protected), "Yes" if hits[SLT[0]].get(tid) else "No"]
        for col, v in enumerate(row, 1):
            ws.cell(row=r, column=col, value=v)
        latest = lambda col: (f'=IF(O{r}=0,"",LOOKUP(2,1/(({ha_id}=A{r})*({ha_date}=P{r})),'
                              f'{ha}!${col}$2:${col}${LAST}))')
        formulas = [
            f"=COUNTIF({ha_id},A{r})",
            f'=IF(O{r}=0,"",_xlfn.MAXIFS({ha_date},{ha_id},A{r}))',
            latest("E"), latest("G"), latest("H"),
            f"=COUNTIF({ml_id},A{r})",
            f'=IF(T{r}=0,"",_xlfn.MAXIFS({ml_date},{ml_id},A{r}))',
            f'=COUNTIFS({ml_id},A{r},{ml_status},"<>Completed",{ml_status},"<>Cancelled")',
        ]
        for col, f in enumerate(formulas, 15):
            c = ws.cell(row=r, column=col, value=f)
            c.fill = FORMULA_FILL
        ws.cell(row=r, column=16).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=21).number_format = "yyyy-mm-dd"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(trees) + 1}"
    ws["M1"].comment = Comment("Area_IDs from the Protected_Areas sheet whose polygon contains this tree "
                               "(point-in-polygon test run when the workbook was built).", "build_workbook.py")
    ws["O1"].comment = Comment("Green columns are formulas linking this tree to the Health_Assessments and "
                               "Maintenance_Log sheets by Tree_ID. Copy them down for trees you add.",
                               "build_workbook.py")

    # ---- Health_Assessments ---------------------------------------------
    ws = wb.create_sheet(ha)
    headers = ["Assessment_ID", "Tree_ID", "Assessment_Date", "Assessor", "Condition", "Issue_Type",
               "Risk_Rating", "Recommended_Outcome", "Council_Request_Ref", "Notes"]
    style_header(ws, headers, [14, 12, 14, 18, 11, 30, 11, 34, 18, 50])
    example = ["HA-0001", "SHW-0661", today, "EXAMPLE – replace or delete this row", "Fair",
               "Fallen or broken branches", "Medium", "Pruning", "", "Example: broken limb over path, crown otherwise healthy."]
    for col, v in enumerate(example, 1):
        c = ws.cell(row=2, column=col, value=v)
        c.font = EXAMPLE_FONT
    ws["C2"].number_format = "yyyy-mm-dd"
    for col in range(1, len(headers) + 1):
        for r in range(3, 52):
            ws.cell(row=r, column=col).fill = INPUT_FILL
    for r in range(2, 52):
        ws.cell(row=r, column=3).number_format = "yyyy-mm-dd"
    add_list_validation(ws, "E", "=Lookups!$A$2:$A$6", "Tree condition")
    add_list_validation(ws, "F", "=Lookups!$B$2:$B$12", "Issue observed (Council request reasons)")
    add_list_validation(ws, "G", "=Lookups!$C$2:$C$5", "Risk rating")
    add_list_validation(ws, "H", "=Lookups!$D$2:$D$5", "Council assessment outcome")

    # ---- Maintenance_Log -------------------------------------------------
    ws = wb.create_sheet(ml)
    headers = ["Work_ID", "Tree_ID", "Assessment_ID", "Work_Date", "Work_Type", "Clearance_Context",
               "Required_Clearance_m", "Status", "Crew_or_Contractor", "Council_Request_Ref", "Notes"]
    style_header(ws, headers, [12, 12, 14, 12, 32, 28, 12, 12, 20, 18, 50],
                 {7: PatternFill("solid", fgColor="548235")})
    example = ["WK-0001", "SHW-0661", "HA-0001", today, "Clearance pruning", "Verge or footpath", None,
               "Scheduled", "", "", "EXAMPLE – replace or delete this row"]
    for col, v in enumerate(example, 1):
        ws.cell(row=2, column=col, value=v).font = EXAMPLE_FONT
    for r in range(2, LAST + 1):
        c = ws.cell(row=r, column=7, value=f'=IF(F{r}="","",IFERROR(INDEX(Lookups!$G$2:$G$6,'
                                            f'MATCH(F{r},Lookups!$F$2:$F$6,0)),""))')
        c.fill = FORMULA_FILL
        c.number_format = "0.0"
    for r in range(2, 52):
        ws.cell(row=r, column=4).number_format = "yyyy-mm-dd"
        if r > 2:
            for col in (1, 2, 3, 4, 5, 6, 8, 9, 10, 11):
                ws.cell(row=r, column=col).fill = INPUT_FILL
    add_list_validation(ws, "E", "=Lookups!$E$2:$E$11", "Type of work done")
    add_list_validation(ws, "F", "=Lookups!$F$2:$F$6", "Where the tree is (sets the clearance)")
    add_list_validation(ws, "H", "=Lookups!$H$2:$H$6", "Job status")

    # ---- Protected_Areas ------------------------------------------------
    ws = wb.create_sheet("Protected_Areas")
    headers = ["Area_ID", "Layer", "Category", "Description", "Centroid_Latitude", "Centroid_Longitude", "Area_m2"]
    style_header(ws, headers, [16, 40, 52, 30, 14, 14, 12])
    for r, a in enumerate(areas, 2):
        for col, v in enumerate(a, 1):
            ws.cell(row=r, column=col, value=v)
        ws.cell(row=r, column=7).number_format = "#,##0"
    ws.auto_filter.ref = f"A1:G{len(areas) + 1}"

    # ---- Lookups ---------------------------------------------------------
    ws = wb.create_sheet("Lookups")
    headers = ["Condition", "Issue_Type", "Risk_Rating", "Recommended_Outcome", "Work_Type",
               "Clearance_Context", "Min_Clearance_m", "Status"]
    style_header(ws, headers, [12, 34, 12, 38, 34, 30, 14, 12])
    ws.freeze_panes = None
    cols = [CONDITIONS, ISSUE_TYPES, RISKS, OUTCOMES, WORK_TYPES, [c for c, _ in CLEARANCES],
            [m for _, m in CLEARANCES], STATUSES]
    for col, values in enumerate(cols, 1):
        for r, v in enumerate(values, 2):
            ws.cell(row=r, column=col, value=v)
    ws["B1"].comment = Comment("Council's tree maintenance and removal request reasons, plus 'None observed', "
                               "'Pests', 'Storm damage' and 'Other'.", "build_workbook.py")
    ws["D1"].comment = Comment("Council's four published assessment outcomes (Brisbane trees page).",
                               "build_workbook.py")
    ws["A1"].comment = Comment("Assumption: Council publishes no condition scale; this one is a placeholder "
                               "to replace with your own standard.", "build_workbook.py")
    ws["C1"].comment = Comment("Assumption: Council publishes no risk scale; replace with your own "
                               "(e.g. a QTRA or TRAQ-based rating).", "build_workbook.py")
    ws["G1"].comment = Comment("Source: Council minimum low-branch clearances, Brisbane trees page, "
                               "https://www.brisbane.qld.gov.au/environment-and-water/plants--trees-and-gardening/brisbane-trees",
                               "build_workbook.py")

    # ---- Sources ---------------------------------------------------------
    ws = wb.create_sheet("Sources")
    style_header(ws, ["Dataset", "Used for", "Records", "URL"], [48, 46, 10, 110])
    ws.freeze_panes = None
    src = [
        ("Botanic collection — Sherwood Arboretum", "Tree_Register (individual trees with location)",
         len(trees), f"{ARCGIS}/{SHERWOOD}/FeatureServer/0"),
    ]
    for code, name, svc in NALL_LAYERS:
        src.append((f"Protected Vegetation (NALL 2003) — {name}", "Protected_Areas; register column M",
                    sum(1 for a in areas if a[0].startswith(code + "-")), f"{ARCGIS}/{svc}/FeatureServer/0"))
    src += [
        (SLT[1], "Protected_Areas; register column N", sum(1 for a in areas if a[0].startswith("CP-SLT-")),
         f"{ARCGIS}/{SLT[2]}/FeatureServer/0"),
        ("Brisbane trees (Council web page)", "Lookups: outcomes, clearances, request reasons", None,
         "https://www.brisbane.qld.gov.au/environment-and-water/plants--trees-and-gardening/brisbane-trees"),
        ("Request tree removal assessment (Council web page)", "Lookups: removal reasons", None,
         "https://www.brisbane.qld.gov.au/online-services/report-an-issue/report-it-tree-removal"),
    ]
    for r, row in enumerate(src, 2):
        for col, v in enumerate(row, 1):
            ws.cell(row=r, column=col, value=v)

    # ---- README ----------------------------------------------------------
    readme.column_dimensions["A"].width = 120
    lines = [
        ("Brisbane tree maintenance register", Font(name=FONT, bold=True, size=14)),
        (f"Built {today.isoformat()} by build_workbook.py from Brisbane City Council open data.", None),
        ("", None),
        ("How the sheets link", Font(name=FONT, bold=True)),
        ("Tree_Register: one row per tree, keyed by Tree_ID (e.g. SHW-0661). Columns A–N are Council data; "
         "green columns O–V are formulas that summarise that tree's rows in the two log sheets.", None),
        ("Health_Assessments: one row per inspection. Enter the Tree_ID; pick Condition, Issue_Type, Risk_Rating "
         "and Recommended_Outcome from the dropdowns.", None),
        ("Maintenance_Log: one row per job. Enter the Tree_ID (and Assessment_ID if the job came from an "
         "inspection). Required_Clearance_m fills in from Clearance_Context.", None),
        ("Protected_Areas: every Natural Assets Local Law protected-vegetation polygon and Significant landscape "
         "tree overlay site, with its centroid. Register column M lists the protected areas each tree falls in.",
         None),
        ("Lookups: dropdown values. Outcomes, request reasons and clearances come from Council's pages; the "
         "condition and risk scales are placeholders (see the cell comments).", None),
        ("", None),
        ("What to edit", Font(name=FONT, bold=True)),
        ("Yellow cells in Health_Assessments and Maintenance_Log are for your entries (any row down to 5001 works; "
         "only the first 50 are shaded). Green cells are formulas: don't type over them.", None),
        ("Grey italic row 2 in each log is an example showing the format. Replace or delete it.", None),
        ("To add trees you survey yourself, add rows to Tree_Register with your own Tree_ID prefix and copy the "
         "green formulas down.", None),
        ("", None),
        ("Limitations of the source data", Font(name=FONT, bold=True)),
        ("Individual tree locations: Council publishes them only for the Sherwood Arboretum (1,181 trees). There "
         "is no public citywide street or park tree inventory, so most of Brisbane's trees are not in this register.",
         None),
        ("Health and maintenance: Council publishes its process (request reasons, arborist assessment, four "
         "outcomes, clearance rules) but no per-tree health or works records. The two log sheets start empty "
         "apart from the example rows; fill them from your own inspections, or from Council service-request "
         "data if you obtain it.", None),
        ("Protected_Areas rows are areas of vegetation, not trees. A polygon may hold one tree or a whole bushland "
         "patch, and one place can sit in more than one layer.", None),
        ("", None),
        ("Refreshing", Font(name=FONT, bold=True)),
        ("Re-run: python3 build_workbook.py brisbane_tree_maintenance.xlsx. This replaces the whole file, "
         "including both logs, so copy your log rows out first and paste them back afterwards.", None),
    ]
    for r, (text, font) in enumerate(lines, 1):
        c = readme.cell(row=r, column=1, value=text)
        c.font = font or Font(name=FONT)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if c.value is not None and c.row > 1 and c.font.name != FONT:
                    f = c.font
                    c.font = Font(name=FONT, bold=f.bold, italic=f.italic, color=f.color, size=f.size)
    wb.save(out_path)
    print(f"{out_path}: {len(trees)} trees, {len(areas)} areas, "
          f"{sum(1 for t in trees if any(hits[c].get(t['Tree_ID']) for c, _, _ in NALL_LAYERS))} trees in NALL areas")


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "brisbane_tree_maintenance.xlsx")
